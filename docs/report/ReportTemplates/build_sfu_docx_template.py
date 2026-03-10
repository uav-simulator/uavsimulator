#!/usr/bin/env python3
"""
Builds a reusable SFU DOCX template according to STU 7.5-07-2021.

Generated files:
  - docs/report/practice/sfu_academic_template.docx
  - docs/report/practice/sfu_report_template.docx (compat alias)
"""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUTPUT_MAIN = Path("docs/report/practice/sfu_academic_template.docx")
OUTPUT_ALIAS = Path("docs/report/practice/sfu_report_template.docx")


def set_font(style, *, size=14, bold=False, italic=False):
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic


def ensure_p_style(doc: Document, name: str, base: str = "Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        return style


def add_page_field(paragraph):
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_page(doc: Document):
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)

    # STU 7.5-07-2021, п. 7.1.2
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(1.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    # Page number at bottom center; hidden on first page.
    sec.different_first_page_header_footer = True
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_page_field(p)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    set_font(normal, size=14)
    n = normal.paragraph_format
    n.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    n.first_line_indent = Cm(1.25)
    n.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    n.space_before = Pt(0)
    n.space_after = Pt(0)

    body = ensure_p_style(doc, "SFU Body")
    set_font(body, size=14)
    p = body.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.first_line_indent = Cm(1.25)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(0)

    title_center = ensure_p_style(doc, "SFU Title Center")
    set_font(title_center, size=14)
    p = title_center.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.first_line_indent = Cm(0)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(0)

    title_small = ensure_p_style(doc, "SFU Title Small")
    set_font(title_small, size=12)
    p = title_small.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.first_line_indent = Cm(0)
    p.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.space_before = Pt(0)
    p.space_after = Pt(0)

    structural = ensure_p_style(doc, "SFU Structural Heading")
    set_font(structural, size=14, bold=True)
    p = structural.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.first_line_indent = Cm(0)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(12)

    h1 = ensure_p_style(doc, "SFU Heading 1")
    set_font(h1, size=14, bold=True)
    p = h1.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.first_line_indent = Cm(1.25)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(12)

    h2 = ensure_p_style(doc, "SFU Heading 2")
    set_font(h2, size=14, bold=True)
    p = h2.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.first_line_indent = Cm(1.25)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(12)

    table_cap = ensure_p_style(doc, "SFU Table Caption")
    set_font(table_cap, size=14)
    p = table_cap.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.first_line_indent = Cm(0)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(12)
    p.space_after = Pt(6)

    fig_cap = ensure_p_style(doc, "SFU Figure Caption")
    set_font(fig_cap, size=14)
    p = fig_cap.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.first_line_indent = Cm(0)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(6)
    p.space_after = Pt(12)

    list_dash = ensure_p_style(doc, "SFU List Dash")
    set_font(list_dash, size=14)
    p = list_dash.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.first_line_indent = Cm(-0.63)
    p.left_indent = Cm(1.88)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(0)

    biblio = ensure_p_style(doc, "SFU Bibliography")
    set_font(biblio, size=14)
    p = biblio.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.first_line_indent = Cm(1.25)
    p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.space_before = Pt(0)
    p.space_after = Pt(0)


def add_blank_lines(doc: Document, count: int):
    for _ in range(count):
        doc.add_paragraph("", style="SFU Body")


def add_title_page(doc: Document):
    # Based on STU 7.5-07-2021, Appendix K (title form for practice reports).
    header_lines = [
        "Министерство науки и высшего образования РФ",
        "Федеральное государственное автономное",
        "образовательное учреждение высшего образования",
        "«СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»",
        "______________________________________________",
        "институт",
        "______________________________________________",
        "кафедра",
    ]
    for line in header_lines:
        doc.add_paragraph(line, style="SFU Title Center")

    add_blank_lines(doc, 4)

    p = doc.add_paragraph("ОТЧЕТ О ПРАКТИКЕ", style="SFU Title Center")
    for run in p.runs:
        run.bold = True

    doc.add_paragraph("______________________________________________", style="SFU Title Center")
    doc.add_paragraph("место прохождения практики", style="SFU Title Small")
    doc.add_paragraph("______________________________________________", style="SFU Title Center")
    doc.add_paragraph("тема", style="SFU Title Small")

    add_blank_lines(doc, 2)

    # Signature block (cleaner visually than underscores in one line).
    t = doc.add_table(rows=6, cols=2)
    t.autofit = True
    left = [
        "Руководитель от университета",
        "",
        "Руководитель от предприятия",
        "",
        "Студент",
        "",
    ]
    right = [
        "________  ____________________",
        "подпись, дата     инициалы, фамилия",
        "________  ____________________",
        "подпись, дата     инициалы, фамилия",
        "________  ____________________",
        "подпись, дата     инициалы, фамилия",
    ]
    for i in range(6):
        t.cell(i, 0).text = left[i]
        t.cell(i, 1).text = right[i]

    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.style = doc.styles["SFU Title Center"]
                if p.text.startswith("подпись"):
                    p.style = doc.styles["SFU Title Small"]

    # Remove table borders for title page signature block.
    tbl = t._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        tblBorders.append(elem)
    tblPr.append(tblBorders)

    add_blank_lines(doc, 3)
    doc.add_paragraph("Красноярск 20__", style="SFU Title Center")
    doc.add_page_break()


def add_clean_skeleton(doc: Document):
    # Minimal skeleton that users can keep or remove.
    doc.add_paragraph("СОДЕРЖАНИЕ", style="SFU Structural Heading")
    doc.add_paragraph("Введение", style="SFU Body")
    doc.add_paragraph("1 Наименование раздела", style="SFU Body")
    doc.add_paragraph("Заключение", style="SFU Body")
    doc.add_paragraph("Список использованных источников", style="SFU Body")
    doc.add_page_break()

    doc.add_paragraph("ВВЕДЕНИЕ", style="SFU Structural Heading")
    doc.add_paragraph("Текст введения.", style="SFU Body")

    doc.add_paragraph("1 Наименование раздела", style="SFU Heading 1")
    doc.add_paragraph("Текст раздела.", style="SFU Body")

    doc.add_paragraph("1.1 Наименование подраздела", style="SFU Heading 2")
    doc.add_paragraph("Текст подраздела.", style="SFU Body")

    doc.add_paragraph("- Первый пункт перечисления", style="SFU List Dash")
    doc.add_paragraph("- Второй пункт перечисления", style="SFU List Dash")

    doc.add_paragraph("Таблица 1 - Наименование таблицы", style="SFU Table Caption")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Параметр"
    table.cell(0, 1).text = "Значение"
    table.cell(0, 2).text = "Примечание"
    table.cell(1, 0).text = "Пример"
    table.cell(1, 1).text = "0"
    table.cell(1, 2).text = "-"

    doc.add_paragraph("Рисунок 1 - Наименование рисунка", style="SFU Figure Caption")

    doc.add_paragraph("ЗАКЛЮЧЕНИЕ", style="SFU Structural Heading")
    doc.add_paragraph("Текст заключения.", style="SFU Body")

    doc.add_paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", style="SFU Structural Heading")
    doc.add_paragraph("[1] Пример библиографической записи.", style="SFU Bibliography")

    doc.add_paragraph("ПРИЛОЖЕНИЕ А", style="SFU Structural Heading")
    doc.add_paragraph("Наименование приложения", style="SFU Structural Heading")
    doc.add_paragraph("Содержание приложения.", style="SFU Body")


def build_template(path: Path):
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    add_title_page(doc)
    add_clean_skeleton(doc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main():
    build_template(OUTPUT_MAIN)
    # Keep old filename for compatibility with previous references.
    build_template(OUTPUT_ALIAS)
    print(f"Template saved: {OUTPUT_MAIN}")
    print(f"Template alias saved: {OUTPUT_ALIAS}")


if __name__ == "__main__":
    main()
