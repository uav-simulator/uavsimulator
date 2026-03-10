#!/usr/bin/env python3
"""
Creates a reusable DOCX template using styles/layout from the already defended
report `Отчет по практике 2024.docx` without modifying the source file.
"""

from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SOURCE = Path("docs/report/practice/Отчет по практике 2024.docx")
TARGET = Path("docs/report/practice/sfu_template_from_2024.docx")


def clear_body_keep_section(doc: Document):
    body = doc._body._element
    for child in list(body):
        # Keep section properties (page setup/footer linkage).
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_title_line(doc: Document, text: str):
    p = doc.add_paragraph(text, style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = 0
    return p


def add_blank(doc: Document, count: int):
    for _ in range(count):
        p = doc.add_paragraph("", style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def build_template():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source file not found: {SOURCE}")

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, TARGET)

    doc = Document(TARGET)
    clear_body_keep_section(doc)

    # Title page based on actual defended report formatting.
    add_title_line(
        doc,
        "Министерство науки и высшего образования РФ Федеральное государственное автономное образовательное учреждение высшего образования",
    )
    add_title_line(doc, "«СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»")
    add_blank(doc, 1)
    add_title_line(doc, "Институт космических и информационных технологий")
    add_title_line(doc, "институт")
    add_title_line(doc, "Кафедра «Программная инженерия»")
    add_title_line(doc, "кафедра")

    add_blank(doc, 6)
    add_title_line(doc, "ОТЧЕТ ОБ УЧЕБНОЙ ПРАКТИКЕ")
    add_title_line(doc, "______________________________________________")
    add_title_line(doc, "место прохождения практики")
    add_title_line(doc, "______________________________________________")
    add_title_line(doc, "тема")
    add_blank(doc, 8)
    add_title_line(doc, "Красноярск 20__")

    doc.add_page_break()

    # Body skeleton.
    doc.add_paragraph("ВВЕДЕНИЕ", style="Heading 3")
    doc.add_paragraph("Текст введения.", style="Normal")

    doc.add_paragraph("1 Наименование раздела", style="Heading 1")
    doc.add_paragraph("Текст раздела.", style="Normal")

    doc.add_paragraph("1.1 Наименование подраздела", style="Heading 2")
    doc.add_paragraph("Текст подраздела.", style="Normal")

    doc.add_paragraph("Таблица 1 – Наименование таблицы", style="Caption")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(0, 2).text = "Примечание"
    table.cell(1, 0).text = "Пример"
    table.cell(1, 1).text = "-"
    table.cell(1, 2).text = "-"

    doc.add_paragraph("Рисунок 1 – Наименование рисунка", style="Рисунок")

    doc.add_paragraph("ЗАКЛЮЧЕНИЕ", style="Heading 3")
    doc.add_paragraph("Итоговые выводы по работе.", style="Normal")

    doc.add_paragraph("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", style="Heading 3")
    doc.add_paragraph("[1] Пример библиографического описания.", style="Normal")

    doc.add_paragraph("ПРИЛОЖЕНИЕ А", style="Heading 3")
    doc.add_paragraph("Материалы приложения.", style="Normal")

    doc.save(TARGET)
    print(f"Template saved: {TARGET}")


if __name__ == "__main__":
    build_template()
