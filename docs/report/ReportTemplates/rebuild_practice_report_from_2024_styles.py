#!/usr/bin/env python3
"""
Rebuilds practice_report.docx using styles/layout from the defended report 2024.
Source style donor is NOT modified.
"""

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

SOURCE_STYLES = Path("docs/report/practice/Отчет по практике 2024.docx")
TARGET_REPORT = Path("docs/report/practice/practice_report.docx")
BACKUP_REPORT = Path("docs/report/practice/practice_report.pre_restyle.bak.docx")
ASSETS_DIR = Path("docs/report/practice/assets")

BODY = [
    ("heading3", "ВВЕДЕНИЕ"),
    (
        "normal",
        "В рамках учебной практики выполнялась прикладная инженерная задача по разработке программного драйвера и пользовательской web-системы для дистанционного управления роботизированной платформой Keyestudio KS0223 на базе Raspberry Pi.",
    ),
    (
        "normal",
        "Практическая значимость работы заключается в построении целостного контура: от передачи управляющих команд до визуализации телеметрии и стабилизации работы системы после перезапуска оборудования.",
    ),
    ("heading1", "ЦЕЛЬ И ЗАДАЧИ ПРАКТИКИ"),
    (
        "normal",
        "Целью практики являлось создание работоспособного программного комплекса для управления KS0223 с персонального компьютера, с поддержкой мониторинга состояния платформы и эксплуатационной документацией.",
    ),
    (
        "normal",
        "Для достижения цели были поставлены задачи: изучить действующий протокол обмена с роботом, реализовать backend-адаптер команд, разработать frontend-пульт управления, интегрировать телеметрию и обеспечить устойчивую работу системы после reboot Raspberry Pi.",
    ),
    ("heading1", "ИСХОДНЫЕ УСЛОВИЯ И АНАЛИЗ"),
    (
        "normal",
        "Исходная конфигурация включала существующий TCP-сервер управления на Raspberry Pi (порт 5051), скрипт передачи кадров с камеры и дополнительные сервисы датчиков. На этапе анализа были выявлены особенности протокола управления и ограничения, связанные с отсутствием унифицированного API для отдельных узлов платформы.",
    ),
    (
        "normal",
        "Также была проведена диагностика сетевого поведения платы после перезапуска. Зафиксированы случаи, когда видеопоток продолжал поступать, но канал управления по TCP становился недоступным.",
    ),
    ("heading1", "ВЫПОЛНЕННЫЕ РАБОТЫ"),
    (
        "normal",
        "Разработан backend-сервис на ASP.NET Core (.NET 8), который выступает промежуточным слоем между web-клиентом и роботом. Реализованы HTTP API, SignalR-канал, отправка команд управления, обработка входящих данных и журналирование событий в формате JSONL.",
    ),
    (
        "normal",
        "Разработан frontend на React + TypeScript (Vite, MUI) с интерфейсом оператора: подключение к роботу, управление движением и камерой, аварийная остановка, настройка скорости команд, отображение статусов и ошибок соединения.",
    ),
    (
        "normal",
        "Интегрированы функции работы с датчиками и исполнительными узлами: управление ультразвуковым модулем HC-SR04 (включая сервопривод), вывод телеметрии, настройки overlay для видеопотока, а также управление LED-панелью.",
    ),
    (
        "normal",
        "Подготовлен контейнеризированный запуск проекта (Dockerfile, Makefile, healthcheck), что позволило стандартизировать локальный запуск и обновление компонентов.",
    ),
    ("heading1", "АРХИТЕКТУРА ПРИЛОЖЕНИЯ"),
    (
        "normal",
        "Архитектура решения построена по схеме «web-клиент - локальный backend-адаптер - робот KS0223». Пользователь работает в браузере с React SPA, backend на ASP.NET Core запускается локально на Mac и принимает HTTP/SignalR-запросы от интерфейса, а затем взаимодействует с Raspberry Pi по TCP, UDP и HTTP в пределах уже существующих сетевых каналов робота.",
    ),
    (
        "normal",
        "В backend выделены несколько специализированных служб. PiTcpClientService отвечает за подключение к TCP-порту 5051, повторные попытки соединения, отправку команд и fail-safe STOP. CameraStreamService принимает UDP-кадры и проверяет типовые HTTP camera endpoints. SensorBridgeService взаимодействует с дополнительным bridge API для расширенной телеметрии, управления HC-SR04 и LED-панелью. SessionLogger сохраняет исходящие команды и входящие сообщения в JSONL-журнал.",
    ),
    (
        "normal",
        "Frontend разделён на модули отображения и управления: ConnectionCard, ControlPad, CameraPanel, TelemetryPanel, LogPanel и отдельные вкладки для сенсоров и LED-панели. Важная особенность решения состоит в том, что операторский интерфейс не работает напрямую с Raspberry Pi: все запросы проходят через backend, что упрощает логику клиента и позволяет централизованно обрабатывать ошибки, healthcheck и аварийную остановку.",
    ),
    (
        "normal",
        "К преимуществам выбранной архитектуры относятся: изоляция протокола робота внутри backend-адаптера; независимость frontend от деталей TCP-обмена; наличие realtime-обновлений через SignalR; возможность логирования и диагностики на стороне Mac; расширяемость за счёт фоновых служб и дополнительных bridge-endpoint. К ограничениям относятся зависимость от нескольких сетевых каналов робота одновременно, ограниченная телеметрия по штатному TCP-протоколу и необходимость учитывать сетевые особенности Raspberry Pi после reboot.",
    ),
    (
        "figure",
        str(ASSETS_DIR / "architecture_diagram.png"),
        "Рисунок 1 - Схема архитектуры приложения управления KS0223",
        16.0,
    ),
    (
        "figure",
        str(ASSETS_DIR / "sequence_diagram.png"),
        "Рисунок 2 - Диаграмма последовательности сценария подключения и отправки команд",
        16.0,
    ),
    (
        "figure",
        str(ASSETS_DIR / "ks0223_robot_photo.jpg"),
        "Рисунок 3 - Внешний вид роботизированной платформы Keyestudio KS0223 (временное изображение)",
        14.5,
    ),
    (
        "figure",
        str(ASSETS_DIR / "product_control_page.png"),
        "Рисунок 4 - Основной экран разработанного web-приложения управления KS0223",
        16.0,
    ),
    (
        "figure",
        str(ASSETS_DIR / "product_logs_page.png"),
        "Рисунок 5 - Экран журналирования и управления логами",
        16.0,
    ),
    ("heading1", "УСТРАНЕНИЕ КРИТИЧЕСКОЙ ПРОБЛЕМЫ ПОСЛЕ REBOOT"),
    (
        "normal",
        "Ключевой эксплуатационной проблемой являлась нестабильность сетевого доступа к каналу управления после перезапуска Raspberry Pi. В процессе диагностики обнаружен конфликт сетевых механизмов: одновременно использовались dhcpcd и networking.service с legacy static-профилем eth0.",
    ),
    (
        "normal",
        "Для устранения проблемы подготовлен и внедрён эксплуатационный патч: отключён конфликтный static-профиль /etc/network/interfaces.d/eth0, отключён networking.service, сохранён единый менеджер сети dhcpcd, оставлены приоритеты маршрутов для wlan0/eth0, добавлен сценарий безопасного отката.",
    ),
    (
        "normal",
        "Патч документирован в репозитории и проверен на реальном reboot. По результатам проверки подтверждена стабильная доступность порта 5051, корректная работа сервисов управления и восстановление телеметрии.",
    ),
    ("heading1", "РЕЗУЛЬТАТЫ ПРАКТИКИ"),
    (
        "normal",
        "По итогам практики получен рабочий программный комплекс управления KS0223, включающий серверную и клиентскую части, поддержку телеметрии, средства диагностики и инструкции по сопровождению.",
    ),
    (
        "normal",
        "Решение обеспечивает: подключение к роботу по TCP, потоковое отображение камеры, управление ключевыми узлами платформы, логирование сессий и стабильный сценарий запуска после перезагрузки оборудования.",
    ),
    ("heading3", "ЗАКЛЮЧЕНИЕ"),
    (
        "normal",
        "Поставленные цели и задачи учебной практики выполнены в полном объёме. Реализованный программный контур применим для учебных и исследовательских задач в области робототехники, а также может служить базой для дальнейшего развития: расширения набора сенсоров, автоматизации сценариев движения и интеграции с алгоритмами интеллектуального управления.",
    ),
    ("heading3", "ПРИЛОЖЕНИЕ А. КЛЮЧЕВЫЕ ТЕХНОЛОГИИ"),
    (
        "normal",
        "ASP.NET Core (.NET 8), SignalR, React, TypeScript, Vite, MUI, Docker, Makefile, TCP/UDP, JSONL-логирование, systemd.",
    ),
    ("heading3", "ПРИЛОЖЕНИЕ Б. КОМПОНЕНТЫ И НАБОР ПЛАТФОРМЫ"),
    (
        "normal",
        "На рисунке 6 приведён пример состава набора KS0223 по официальной документации производителя. Изображение используется как временная иллюстрация и может быть заменено на фотографии фактической установки.",
    ),
    (
        "figure",
        str(ASSETS_DIR / "ks0223_kit_list.png"),
        "Рисунок 6 - Пример состава набора KS0223 по документации Keyestudio (временное изображение)",
        10.0,
    ),
    ("heading3", "ПРИЛОЖЕНИЕ В. ФРАГМЕНТЫ КОДА"),
    (
        "listing",
        "Листинг 1 - Регистрация backend-сервисов и healthcheck API",
        """builder.Services.AddSignalR();
builder.Services.AddSingleton<PiTcpClientService>();
builder.Services.AddSingleton<CameraStreamService>();
builder.Services.AddSingleton<SensorBridgeService>();
builder.Services.AddHostedService(serviceProvider => serviceProvider.GetRequiredService<PiTcpClientService>());

app.MapGet("/api/health", (PiTcpClientService service, CameraStreamService cameraService, SensorBridgeService sensorBridgeService) =>
{
    var control = service.GetStatus();
    var camera = cameraService.GetStatus();
    var sensors = sensorBridgeService.GetStatus();
    var status = control.DesiredConnection && !control.TcpConnected ? "degraded" : "ok";
    return Results.Ok(new HealthDto(status, DateTimeOffset.UtcNow, version, control, camera, sensors));
});""",
    ),
    (
        "listing",
        "Листинг 2 - Fail-safe STOP при потере клиента или завершении приложения",
        """public async Task UnregisterUiConnectionAsync(string connectionId)
{
    lock (stateLock)
    {
        uiConnectedClients = Math.Max(0, uiConnectedClients - 1);
        mustStop = uiConnectedClients == 0;
    }

    if (mustStop)
    {
        await TrySendStopBestEffortAsync("last-ui-disconnected", CancellationToken.None);
    }
}

public override async Task StopAsync(CancellationToken cancellationToken)
{
    await TrySendStopBestEffortAsync("app-shutdown", cancellationToken);
    await CloseConnectionAsync(cancellationToken);
}""",
    ),
    (
        "listing",
        "Листинг 3 - Хранение и применение настроек overlay на frontend",
        """const OVERLAY_STORAGE_KEY = 'ks0223_camera_overlay_settings_v2'

function loadOverlaySettings(): OverlaySettings {
  const raw = window.localStorage.getItem(OVERLAY_STORAGE_KEY)
  if (!raw) {
    return defaultOverlaySettings
  }

  const parsed = JSON.parse(raw) as Partial<OverlaySettings>
  return { ...defaultOverlaySettings, ...parsed }
}

useEffect(() => {
  window.localStorage.setItem(OVERLAY_STORAGE_KEY, JSON.stringify(overlay))
}, [overlay])""",
    ),
    (
        "listing",
        "Листинг 4 - Эксплуатационный сетевой патч для устойчивой работы после reboot",
        """# Prevent dual-address conflict on eth0:
if [[ -f /etc/network/interfaces.d/eth0 ]]; then
  sudo_cmd mv /etc/network/interfaces.d/eth0 /etc/network/interfaces.d/eth0.ks0223-disabled
fi
sudo_cmd systemctl disable --now networking.service || true

cat <<'DHCPCD' | sudo_cmd tee -a /etc/dhcpcd.conf >/dev/null
interface wlan0
metric 100

interface eth0
metric 400
DHCPCD""",
    ),
    ("heading3", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"),
    (
        "source",
        "Keyestudio. KS0223 Smart Small Turtle Robot Car V2.0 for Raspberry Pi : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/ (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Keyestudio. KS0223 Description : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/Description.html (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Keyestudio. KS0223 Kit List : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/Kit%20List.html (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Keyestudio. KS0223 Assembly Tutorial : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/1.Assembley%20Tutorial/1.Assembly%20Tutorial.html (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Microsoft Learn. ASP.NET Core SignalR overview : [сайт]. - URL: https://learn.microsoft.com/aspnet/core/signalr/introduction?view=aspnetcore-8.0 (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Microsoft Learn. Minimal APIs overview : [сайт]. - URL: https://learn.microsoft.com/aspnet/core/fundamentals/minimal-apis/overview?view=aspnetcore-8.0 (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "MUI. Material UI - React components : [сайт]. - URL: https://mui.com/material-ui/getting-started/ (дата обращения: 10.03.2026).",
    ),
    (
        "source",
        "Vite. Getting Started Guide : [сайт]. - URL: https://vite.dev/guide/ (дата обращения: 10.03.2026).",
    ),
]


def clear_body_keep_section(doc: Document):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clone_paragraph_format(target, source):
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def set_run_font(run, *, bold=None, underline=None, size_pt=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if bold is not None:
        run.bold = bold
    if underline is not None:
        run.underline = underline
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def add_template_paragraph(doc: Document, donor: Document, donor_idx: int, segments):
    p = doc.add_paragraph(style="Normal")
    clone_paragraph_format(p, donor.paragraphs[donor_idx])
    for segment in segments:
        if isinstance(segment, str):
            text = segment
            options = {}
        else:
            text, options = segment
        run = p.add_run(text)
        set_run_font(run, **options)
    return p


def add_blank_from_donor(doc: Document, donor: Document, donor_idx: int):
    p = doc.add_paragraph(style="Normal")
    clone_paragraph_format(p, donor.paragraphs[donor_idx])
    return p


def add_normal_paragraph(doc: Document, text: str):
    return doc.add_paragraph(text, style="Normal")


def add_figure(doc: Document, image_path: str, caption: str, width_cm: float):
    path = Path(image_path)
    if not path.exists():
        add_normal_paragraph(doc, f"[Изображение временно недоступно: {path.name}]")
        return

    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = 0
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))

    caption_paragraph = doc.add_paragraph(caption, style="Рисунок")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = 0


def add_listing(doc: Document, title: str, code: str):
    title_paragraph = doc.add_paragraph(style="Normal")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.paragraph_format.first_line_indent = 0
    title_paragraph.paragraph_format.left_indent = 0
    title_paragraph.paragraph_format.right_indent = 0
    title_paragraph.paragraph_format.space_before = Pt(6)
    title_paragraph.paragraph_format.space_after = 0
    title_paragraph.paragraph_format.line_spacing = 1.0
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, bold=False, size_pt=14)
    title_run.font.color.rgb = None

    code_style = "Листинг" if "Листинг" in [style.name for style in doc.styles] else "Normal"
    code_paragraph = doc.add_paragraph(style=code_style)
    code_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code_paragraph.paragraph_format.first_line_indent = 0
    code_paragraph.paragraph_format.left_indent = Cm(0.8)
    code_paragraph.paragraph_format.right_indent = 0
    code_paragraph.paragraph_format.line_spacing = 1.0
    code_paragraph.paragraph_format.space_before = 0
    code_paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(code.splitlines()):
        run = code_paragraph.add_run(line)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(10)
        if index != len(code.splitlines()) - 1:
            run.add_break()


def add_source(doc: Document, text: str):
    style_name = "List Paragraph" if "List Paragraph" in [style.name for style in doc.styles] else "Normal"
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.left_indent = 0
    paragraph.paragraph_format.right_indent = 0
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=14)


def remove_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def add_signature_block(doc: Document):
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)

    widths = (Cm(8.2), Cm(3.2), Cm(4.8))
    rows = [
        ("Руководитель от университета", "____________", "____________________"),
        ("Руководитель от предприятия", "____________", "____________________"),
        ("Студент", "____________", "Н.М. Горовенко"),
    ]

    for row_idx, row_values in enumerate(rows):
        row = table.rows[row_idx]
        for col_idx, value in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.width = widths[col_idx]
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = 0
            paragraph.paragraph_format.left_indent = 0
            paragraph.paragraph_format.right_indent = 0
            paragraph.paragraph_format.space_before = 0
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            set_run_font(run, size_pt=14)

    return table


def build_title_page(doc: Document, donor: Document):
    add_template_paragraph(
        doc,
        donor,
        0,
        [
            "Министерство науки и высшего образования РФ ",
            (
                "Федеральное государственное автономное образовательное учреждение высшего образования",
                {"size_pt": 12},
            ),
        ],
    )
    add_template_paragraph(doc, donor, 1, [("«СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»", {"bold": True})])
    add_blank_from_donor(doc, donor, 2)
    add_template_paragraph(
        doc, donor, 3, [("Институт космических и информационных технологий", {"underline": True})]
    )
    add_template_paragraph(doc, donor, 4, [("институт", {"size_pt": 10})])
    add_template_paragraph(doc, donor, 5, [("Кафедра «Программная инженерия»", {"underline": True})])
    add_template_paragraph(doc, donor, 6, [("кафедра", {"size_pt": 10})])

    for donor_idx in range(7, 18):
        add_blank_from_donor(doc, donor, donor_idx)

    add_template_paragraph(doc, donor, 18, [("ОТЧЕТ ОБ УЧЕБНОЙ ПРАКТИКЕ", {"bold": True})])
    add_template_paragraph(
        doc,
        donor,
        19,
        [("ФГАОУ ВО СФУ, ИКИТ, кафедра «Программная инженерия»", {"underline": True})],
    )
    add_template_paragraph(doc, donor, 20, [("место прохождения практики", {"size_pt": 12})])
    add_template_paragraph(
        doc,
        donor,
        21,
        [("Разработка программного драйвера и web-системы управления", {"underline": True})],
    )
    add_template_paragraph(
        doc,
        donor,
        22,
        [("роботизированной платформой Keyestudio KS0223", {"underline": True})],
    )
    add_template_paragraph(doc, donor, 23, [("тема", {"size_pt": 12})])

    for donor_idx in range(24, 29):
        add_blank_from_donor(doc, donor, donor_idx)

    add_signature_block(doc)

    for donor_idx in range(29, 32):
        add_blank_from_donor(doc, donor, donor_idx)

    add_template_paragraph(doc, donor, 43, ["Красноярск 2026"])


def add_body_item(doc: Document, kind: str, text: str):
    if kind == "heading3":
        return doc.add_paragraph(text, style="Heading 3")
    if kind == "heading1":
        return doc.add_paragraph(text, style="Heading 1")
    if kind == "figure":
        image_path, caption, width_cm = text, "", 15.0
        raise ValueError("Figure items must be handled separately with variable payload.")
    if kind == "listing":
        raise ValueError("Listing items must be handled separately with variable payload.")
    if kind == "source":
        return add_source(doc, text)
    return doc.add_paragraph(text, style="Normal")


def rebuild():
    if not SOURCE_STYLES.exists():
        raise FileNotFoundError(f"Styles source not found: {SOURCE_STYLES}")

    if TARGET_REPORT.exists():
        shutil.copy2(TARGET_REPORT, BACKUP_REPORT)

    shutil.copy2(SOURCE_STYLES, TARGET_REPORT)
    doc = Document(TARGET_REPORT)
    clear_body_keep_section(doc)

    donor = Document(SOURCE_STYLES)
    build_title_page(doc, donor)

    doc.add_page_break()

    for item in BODY:
        kind = item[0]
        if kind in {"heading3", "heading1", "normal", "source"}:
            _, text = item
            add_body_item(doc, kind, text)
            continue
        if kind == "figure":
            _, image_path, caption, width_cm = item
            add_figure(doc, image_path, caption, width_cm)
            continue
        if kind == "listing":
            _, title, code = item
            add_listing(doc, title, code)
            continue
        raise ValueError(f"Unsupported body item kind: {kind}")

    doc.save(TARGET_REPORT)
    print(f"Rebuilt report: {TARGET_REPORT}")
    if BACKUP_REPORT.exists():
        print(f"Backup saved: {BACKUP_REPORT}")


if __name__ == "__main__":
    rebuild()
