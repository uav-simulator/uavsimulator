#!/usr/bin/env python3
"""
Rebuild practice_report_2.docx into an academic internship report while
preserving the existing title page and document styles.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[3]
REPORT_PATH = ROOT / "docs/report/practice/practice_report_2.docx"
ASSETS_DIR = ROOT / "docs/report/practice/assets"

ARCHITECTURE_PNG = ASSETS_DIR / "architecture_diagram.png"
SEQUENCE_PNG = ASSETS_DIR / "sequence_diagram.png"
ROBOT_PNG = ASSETS_DIR / "ks0223_robot_photo.jpg"
CONTROL_UI_PNG = ASSETS_DIR / "product_control_page.png"
LOGS_UI_PNG = ASSETS_DIR / "product_logs_page.png"
KIT_LIST_PNG = ASSETS_DIR / "ks0223_kit_list.png"

SYSTEM_OVERVIEW_PNG = ASSETS_DIR / "system_overview_diagram.png"
SYSTEM_INTERACTION_PNG = ASSETS_DIR / "system_interaction_diagram.png"
API_CONTRACT_PNG = ASSETS_DIR / "runtime_api_contract_diagram.png"
CONNECTION_UI_PNG = ASSETS_DIR / "connection_ui_diagram.png"
UNITY_POPUP_PNG = ASSETS_DIR / "unity_popup_diagram.png"
SHARED_WORLD_TABS_PNG = ASSETS_DIR / "shared_world_tabs_diagram.png"
PARALLEL_RUNTIME_PNG = ASSETS_DIR / "parallel_runtime_diagram.png"


INTRO_TEXT = [
    "Учебная практика проходила в 2026 году на базе ФГАОУ ВО «Сибирский федеральный университет», Института космических и информационных технологий, кафедры «Программная инженерия», в сроки, установленные учебным графиком. Практика была ориентирована на выполнение прикладной инженерной задачи, связанной с разработкой программного обеспечения для роботизированной платформы Keyestudio KS0223.",
    "Предметом практики стала разработка единого WebUI-драйвера и web-системы управления, позволяющих работать как с реальным стендом на базе Raspberry Pi, так и с Unity-симулятором. Такой результат имеет самостоятельную прикладную ценность и одновременно служит частью магистерского проекта, связанного с моделированием и sim-to-real сценариями.",
    "В отчёте основное внимание уделено не только характеристике полученного программного продукта, но и ходу выполнения практики: анализу исходного состояния платформы, этапам работ, личному вкладу студента, использованным технологиям, результатам проверки и приобретённым профессиональным навыкам.",
]

GOAL_TEXT = [
    "Индивидуальное задание практики заключалось в разработке единого WebUI-драйвера и web-системы управления роботизированной платформой Keyestudio KS0223 для двух режимов работы: с физическим стендом и с Unity-симулятором. Дополнительно требовалось подготовить эксплуатационные материалы, подтверждающие работоспособность системы и пригодность результата для дальнейшего развития.",
    "Цель практики состояла в получении практических навыков проектирования, разработки, интеграции и проверки программного обеспечения для робототехнической платформы, а также в создании работоспособного программного комплекса, который объединяет управление, видеопоток, сенсорные данные и диагностические функции в одном интерфейсе.",
    "В качестве критериев достижения цели были приняты: установление соединения с режимом `real-robot`; передача базовых команд движения; получение видеокадра и отображение видеопотока; получение и отображение сенсорных данных; работа с Unity-симулятором через тот же интерфейс; журналирование действий и состояния системы; устойчивое восстановление работоспособности после reboot Raspberry Pi.",
    "Основные задачи практики, действия студента и полученные результаты сведены в таблицу 1.",
]

OBJECT_TEXT = [
    "Объектом практики являлась роботизированная платформа Keyestudio KS0223, построенная на базе Raspberry Pi и снабжённая средствами движения, видеокамерой, ультразвуковым датчиком, сервоприводом, линейными датчиками и световой индикацией. На рисунке 1 приведено изображение платформы, использовавшееся в качестве иллюстрации стенда.",
    "Исходное программное обеспечение платформы включало управляющий скрипт `MainControl.py`, принимающий строковые команды по TCP на порту `5051`, а также отдельный механизм передачи кадров камеры. Такая организация означала, что управление, видеопоток и сенсорные данные были разнесены по разным каналам и не представляли собой единый интерфейс оператора.",
    "Для получения расширенных данных о состоянии платформы требовалось использовать отдельный bridge для сенсоров и исполнительных узлов. В ходе анализа было установлено, что штатного TCP-канала недостаточно для построения полноценного web-интерфейса, так как он не возвращает телеметрию и не обеспечивает централизованного контроля состояния системы.",
    "Параллельно в репозитории уже развивался Unity-симулятор, допускающий внешнее управление через runtime API. Это позволило рассматривать объект практики шире, чем только физический стенд: разрабатываемый интерфейс должен был обеспечивать единый сценарий работы как с реальной платформой, так и с цифровой моделью.",
]

STAGES_INTRO = [
    "Центральной частью практики стало последовательное выполнение этапов разработки и проверки программного комплекса. Важной особенностью работы было то, что каждый этап был связан с личным участием студента в анализе, реализации, интеграции и документировании полученного решения.",
]

STAGE_ITEMS = [
    "На первом этапе студент исследовал штатное программное обеспечение KS0223, структуру сетевого взаимодействия и фактические возможности стандартного канала TCP `5051`, а также проверил работу видеоканала и исходных сервисов на Raspberry Pi.",
    "На втором этапе студент проанализировал сетевые каналы управления, передачи видеоданных и сенсорной информации, после чего определил, какие функции должны быть вынесены в backend-сервис, а какие могут оставаться на стороне платформы.",
    "На третьем этапе студент спроектировал серверную часть системы на ASP.NET Core, определил состав основных endpoint, логику работы с двумя режимами исполнения и подход к маршрутизации команд, телеметрии и состояния соединения.",
    "На четвёртом этапе студент реализовал WebUI-интерфейс оператора, включающий экран подключения, основные элементы управления движением, отображение видеопотока, блок диагностики и средства просмотра журналов.",
    "На пятом этапе студент интегрировал канал сенсоров и исполнительных узлов, обеспечил приём и обработку видеопотока, организовал передачу данных о состоянии платформы и связал эти данные с web-интерфейсом.",
    "На шестом этапе студент подготовил локальный и контейнеризированный запуск системы, настроил эксплуатационные сценарии для запуска, сопровождения и диагностики, а также проверил работоспособность программного комплекса в типовых условиях эксплуатации.",
    "На седьмом этапе студент подключил Unity-симулятор к тому же интерфейсу управления, проверил работу с общей сценой симуляции, выбор управляемого агента и камеры, а также сопоставил сценарии работы с реальным стендом и цифровым двойником.",
    "На восьмом этапе студент выполнил диагностику нестабильной работы после перезагрузки Raspberry Pi, подготовил исправления для сетевого и сервисного контура, проверил восстановление работоспособности и оформил полученные результаты в виде рисунков, таблиц, листингов и отчётных материалов.",
]

STAGES_OUTRO = [
    "Таким образом, практика включала не только разработку отдельных модулей, но и полный цикл инженерной работы: от анализа исходного состояния системы до проверки полученного результата и подготовки отчётной документации.",
]

TECH_TEXT = [
    "Для серверной части был выбран стек ASP.NET Core на платформе .NET 8. Это решение позволило организовать web-сервис с HTTP API, обработкой фоновых задач, средствами журналирования и интеграцией с SignalR для доставки данных в режиме, близком к реальному времени.",
    "Для клиентской части использовались React, TypeScript, Vite и библиотека MUI. Такой набор инструментов обеспечил достаточно быстрый цикл разработки интерфейса и позволил реализовать единый экран управления, не разделяя frontend на отдельные приложения для разных режимов работы.",
    "При интеграции с платформой KS0223 студент использовал сетевые протоколы TCP, UDP и HTTP. Применение нескольких протоколов было обусловлено тем, что команды управления, видеоданные и данные сенсоров изначально передавались по разным каналам и требовали согласованной обработки на стороне сервера.",
    "Unity 6000.1.8f1 использовался как программная среда моделирования и цифровой двойник платформы. Наличие внешнего runtime API позволило подключить симулятор к тому же WebUI и тем самым проверить применимость разрабатываемого интерфейса не только к физическому стенду, но и к сценам моделирования.",
    "Для эксплуатационной части применялись Docker, Makefile, Linux systemd и SSH. Эти инструменты использовались для локального развёртывания, сопровождения сервисов на Raspberry Pi, повторяемого запуска и диагностики системы после изменений и перезапусков.",
    "Выбор перечисленных технологий определялся не стремлением использовать максимальное число инструментов, а необходимостью построить целостный и воспроизводимый контур практической работы. Важным результатом практики стало понимание ограничений каждого из средств и способов их совместного использования в инженерной задаче.",
]

TECH_IMPL_TEXT = [
    "Результатом практики стал программный комплекс, объединяющий web-интерфейс оператора, серверную часть, работу с реальным стендом KS0223 и подключение Unity-симулятора. Общая схема системы и последовательность взаимодействия её компонентов приведены на рисунках 2 и 3.",
    "Серверная часть обеспечивает приём запросов от web-интерфейса, хранение состояния клиентской сессии, маршрутизацию команд, получение телеметрии и работу с видеоданными. В практической реализации использовалась схема адресации через `clientId` и `runtimeMode`, благодаря которой стало возможно изолировать состояние отдельных вкладок браузера и поддерживать независимую работу с физическим стендом и симулятором.",
    "Для режима работы с физическим стендом серверный контур взаимодействует с платформой через TCP `5051`, отдельный канал приёма кадров камеры и bridge сенсоров. Это позволило объединить в одном интерфейсе управление движением, просмотр изображения, получение данных HC-SR04, line tracking, LED и связанных диагностических параметров.",
    "Для режима работы с Unity-симулятором использовался внешний runtime API. В рамках практики было обеспечено подключение к общей сцене симуляции, выбор управляемого агента и камеры, а также работа нескольких клиентских сессий в одной сцене. Схема такого сценария приведена на рисунке 6, а более подробные технические схемы вынесены в приложение Б.",
    "Внешний вид итогового интерфейса оператора показан на рисунках 4 и 5. На основном экране реализованы подключение, управление платформой, просмотр текущего состояния и камеры. Отдельный экран используется для диагностики, просмотра журналов и контроля служебных сообщений.",
    "К числу полученных технических результатов относятся: создание единого web-интерфейса, реализация серверной части для двух режимов работы, интеграция сенсорного канала и видеопотока, подготовка средств журналирования, а также оформление кода и схем, представленных в приложениях Б и В.",
]

VERIFICATION_TEXT = [
    "Проверка работоспособности выполнялась в соответствии с критериями, сформулированными в разделе цели и задач практики. Контроль охватывал как базовые пользовательские сценарии, так и эксплуатационные ситуации, связанные с повторными подключениями и перезапуском Raspberry Pi.",
    "Результаты функциональной проверки представлены в таблице 2. Таблица включает основные действия оператора и подтверждает, что разработанная система закрывает минимально необходимый набор практических задач для реального стенда и Unity-симулятора.",
    "Результаты эксплуатационной проверки представлены в таблице 3. В ней зафиксированы типовые условия запуска и восстановления системы, а также состояние ключевых каналов: TCP `5051`, видеопотока, bridge сенсоров и запуска сервисов.",
    "Проведённая проверка показывает, что цель практики достигнута в качественно-эксплуатационном смысле: система обеспечивает работоспособное управление, наблюдение и диагностику для двух режимов работы. При этом дальнейшее количественное сравнение задержек, интенсивности поступления данных и характеристик sim-to-real сценариев целесообразно выполнять уже в рамках магистерского исследования.",
]

SKILLS_TEXT = [
    "В ходе практики студент приобрёл и закрепил навыки, связанные с анализом существующего программного обеспечения, проектированием web-сервисов, разработкой интерфейсов и интеграцией разнородных программно-аппаратных модулей. Практика позволила перейти от изучения отдельных технологий к их совместному применению в единой инженерной задаче.",
    "Сформированные в ходе практики компетенции имеют прикладной характер и напрямую соотносятся с подготовкой магистранта по направлению «Программная инженерия». Важно, что студент не только реализовал программные компоненты, но и выполнил диагностику ограничений исходной платформы, подготовил эксплуатационные материалы и оформил результаты в форме инженерного отчёта.",
]

SKILL_ITEMS = [
    "анализ существующего программного обеспечения и выявление его функциональных ограничений;",
    "работа с сетевыми протоколами TCP, UDP и HTTP при интеграции распределённых компонентов;",
    "проектирование серверной части web-системы и интерфейсов взаимодействия модулей;",
    "разработка пользовательского web-интерфейса для управления и наблюдения за системой;",
    "интеграция сенсоров, видеопотока и диагностических сообщений в единый программный контур;",
    "отладка сервисов Linux и конфигурации Raspberry Pi в эксплуатационных сценариях;",
    "подготовка технической документации, рисунков, таблиц и приложений по результатам работы.",
]

CONCLUSION_TEXT = [
    "В ходе учебной практики студент выполнил индивидуальное задание по разработке единого WebUI-драйвера и web-системы управления для платформы Keyestudio KS0223. Выполненная работа включала анализ исходного состояния платформы, проектирование и реализацию программных модулей, интеграцию с Unity-симулятором, проверку работоспособности и оформление результата в форме отчёта и приложений.",
    "Поставленная цель достигнута: разработанный программный комплекс обеспечивает подключение к реальному стенду и Unity-симулятору, передачу управляющих команд, получение видеоданных и сенсорной информации, журналирование и базовую эксплуатационную устойчивость. Практика сформировала прикладную основу для дальнейшей магистерской работы, в рамках которой полученный интерфейс может использоваться как операторский слой для sim-to-real сценариев.",
]

APPENDIX_A_TEXT = [
    "В приложении А приведены основные технологии, использованные в ходе практики, и кратко пояснено их назначение.",
    "ASP.NET Core (.NET 8) использовался для реализации серверной части системы, HTTP API, фоновых задач и средств журналирования. Этот стек оказался удобен для построения прикладного сервиса, связывающего web-интерфейс, Raspberry Pi и Unity-симулятор.",
    "SignalR применялся для доставки телеметрии, состояний и служебных сообщений в привязке к клиентской сессии. Его использование позволило упростить передачу данных в интерфейс без постоянного опроса всех endpoint со стороны клиента.",
    "React, TypeScript, Vite и MUI использовались для построения операторского интерфейса. Они обеспечили достаточную гибкость в разработке web-форм подключения, элементов управления и экранов диагностики.",
    "Unity 6000.1.8f1 использовался как среда моделирования, в которой была доступна сцена с роботизированной платформой и внешним runtime API. Это позволило в рамках практики работать не только с физическим стендом, но и с цифровым двойником.",
    "Docker, Makefile, Linux systemd, SSH, TCP, UDP, HTTP и JSONL-журналирование применялись как инфраструктурные средства сопровождения, эксплуатации и диагностики. Их совместное использование позволило сделать программный комплекс воспроизводимым и удобным для проверки в типовых сценариях.",
]

APPENDIX_B_TEXT = [
    "В приложении Б собраны дополнительные схемы, дополняющие основное описание технической реализации. Эти материалы не являются центром отчёта по практике, однако позволяют зафиксировать ключевые инженерные решения, использованные в ходе разработки.",
]

APPENDIX_D_TEXT = [
    "К дальнейшим направлениям развития относятся расширение сравнения поведения системы в режимах `sim` и `real`, углубление количественного анализа эксплуатационных характеристик, подготовка фотографий фактического стенда для финальной версии материалов и развитие исследовательского слоя, связанного с подключением моделей управления и внешних Python/ROS2 инструментов.",
    "На рисунке 13 приведён пример состава набора KS0223 по документации производителя. Эта иллюстрация используется как дополнительный справочный материал и может быть заменена фотографиями конкретной аппаратной конфигурации, применявшейся в ходе демонстрации.",
]

TABLE_1_ROWS = [
    ["Задача практики", "Выполненные действия студента", "Полученный результат"],
    ["Анализ каналов взаимодействия с KS0223 и Unity", "Исследованы `MainControl.py`, `FramesSend.py`, каналы TCP/UDP/HTTP и runtime API Unity", "Определены исходные ограничения платформы и состав интеграционных точек"],
    ["Разработка серверной части", "Спроектирован и реализован backend на ASP.NET Core для двух режимов работы", "Создан серверный контур управления, диагностики и передачи данных"],
    ["Разработка WebUI", "Реализован единый web-интерфейс оператора", "Получен интерфейс для подключения, управления, просмотра камеры и логов"],
    ["Интеграция телеметрии и периферии", "Подключён bridge сенсоров и исполнительных узлов", "В интерфейсе доступны данные датчиков и управляющие действия"],
    ["Организация видеоканала", "Настроен приём и обработка кадров камеры", "Подтверждена возможность получения snapshot и видеопотока"],
    ["Подготовка запуска и сопровождения", "Настроены Docker, Makefile и эксплуатационные сценарии", "Система запускается локально и сопровождается по единым инструкциям"],
    ["Устранение проблемы после reboot", "Проведена диагностика сервисов и сетевого контура Raspberry Pi", "Подтверждено восстановление работоспособности после контрольного reboot"],
    ["Подготовка отчётных материалов", "Собраны рисунки, таблицы, листинги и текст отчёта", "Результаты практики оформлены в академическом виде"],
]

TABLE_2_ROWS = [
    ["Функциональный сценарий", "Способ проверки", "Наблюдаемый результат", "Итог"],
    ["Подключение к `real-robot`", "Через экран подключения и серверный endpoint соединения", "Соединение устанавливается, состояние отображается в интерфейсе", "выполнено"],
    ["Отправка базовых команд движения", "Команды вперёд, назад, повороты и остановка", "Платформа реагирует на управляющие воздействия", "выполнено"],
    ["Получение snapshot", "Запрос snapshot камеры в интерфейсе и через backend", "Формируется и отображается актуальный кадр", "выполнено"],
    ["Приём видеопотока", "Проверка потока кадров от камеры", "Кадры поступают и могут отображаться оператору", "выполнено"],
    ["Получение сенсорных данных", "Чтение данных bridge сенсоров", "В интерфейсе доступны показания датчиков и состояние узлов", "выполнено"],
    ["Просмотр журналов и диагностики", "Открытие экрана журналирования и диагностики", "Доступны записи событий и сообщения о состоянии системы", "выполнено"],
    ["Подключение к Unity-симулятору", "Выбор режима `unity-sim` и подключение к runtime", "WebUI работает с цифровым двойником через тот же интерфейс", "выполнено"],
    ["Выбор `control agent` и `camera agent`", "Настройка параметров клиентской сессии в Unity", "Поддерживается раздельный выбор управляемого агента и камеры", "выполнено"],
]

TABLE_3_ROWS = [
    ["Сценарий проверки", "TCP `5051`", "Камера", "Bridge сенсоров", "Запуск сервисов / среды", "Итог"],
    ["Обычный запуск web-комплекса", "доступен", "доступна", "доступен", "службы и backend запускаются штатно", "успешно"],
    ["Повторное подключение к KS0223", "доступен", "сохраняется доступ", "сохраняется доступ", "сессия восстанавливается без ручной перенастройки", "успешно"],
    ["Контрольный reboot Raspberry Pi", "восстанавливается после перезапуска", "восстанавливается после запуска сервисов", "доступен после старта служб", "подтверждён предсказуемый запуск сервисов", "успешно"],
    ["Подключение к Unity runtime", "не требуется", "виртуальная камера доступна", "телеметрия доступна через runtime", "Unity runtime и WebUI работают совместно", "успешно"],
    ["Параллельная работа нескольких вкладок", "адресуется по клиентской сессии", "камера выбирается для активной сессии", "данные изолируются по сессии", "одновременные сценарии не конфликтуют", "успешно"],
]

SOURCES = [
    "Keyestudio. KS0223 Smart Small Turtle Robot Car for Raspberry Pi : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/ (дата обращения: 13.03.2026).",
    "Keyestudio. KS0223 Description : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/Description.html (дата обращения: 13.03.2026).",
    "Keyestudio. KS0223 Kit List : [сайт]. - URL: https://docs.keyestudio.com/projects/KS0223/en/latest/Kit%20List.html (дата обращения: 13.03.2026).",
    "Microsoft Learn. ASP.NET Core SignalR overview : [сайт]. - URL: https://learn.microsoft.com/aspnet/core/signalr/introduction?view=aspnetcore-8.0 (дата обращения: 13.03.2026).",
    "Microsoft Learn. Hosted services in ASP.NET Core : [сайт]. - URL: https://learn.microsoft.com/aspnet/core/fundamentals/host/hosted-services?view=aspnetcore-8.0 (дата обращения: 13.03.2026).",
    "Microsoft Learn. Minimal APIs overview : [сайт]. - URL: https://learn.microsoft.com/aspnet/core/fundamentals/minimal-apis/overview?view=aspnetcore-8.0 (дата обращения: 13.03.2026).",
    "Unity Manual. Unity 6 Manual : [сайт]. - URL: https://docs.unity3d.com/6000.1/Documentation/Manual/ (дата обращения: 13.03.2026).",
    "MUI. Material UI - React components : [сайт]. - URL: https://mui.com/material-ui/getting-started/ (дата обращения: 13.03.2026).",
    "Vite. Getting Started Guide : [сайт]. - URL: https://vite.dev/guide/ (дата обращения: 13.03.2026).",
    "Docker Docs. Docker overview : [сайт]. - URL: https://docs.docker.com/get-started/docker-overview/ (дата обращения: 13.03.2026).",
    "Open Robotics. ROS 2 Documentation : [сайт]. - URL: https://docs.ros.org/en/humble/index.html (дата обращения: 13.03.2026).",
]


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


FONT_REGULAR = load_font(30, bold=False)
FONT_SMALL = load_font(24, bold=False)
FONT_BOLD = load_font(32, bold=True)
FONT_TITLE = load_font(38, bold=True)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont) -> tuple[int, int]:
    left, top, right, bottom = draw.multiline_textbbox((0, 0), text, font=font, spacing=6, align="center")
    return right - left, bottom - top


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    font: ImageFont.ImageFont,
    fill: str = "#1f2937",
) -> None:
    x1, y1, x2, y2 = box
    width, height = text_size(draw, text, font)
    draw.multiline_text(
        ((x1 + x2 - width) / 2, (y1 + y2 - height) / 2),
        text,
        font=font,
        fill=fill,
        spacing=6,
        align="center",
    )


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    fill: str,
    outline: str = "#111827",
    radius: int = 24,
    font: ImageFont.ImageFont = FONT_REGULAR,
    text_fill: str = "#111827",
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=4)
    draw_centered_text(draw, box, text, font=font, fill=text_fill)


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    fill: str = "#334155",
    width: int = 5,
    head: int = 18,
) -> None:
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=fill, width=width)
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        points = [
            (x2, y2),
            (x2 - direction * head, y2 - head // 2),
            (x2 - direction * head, y2 + head // 2),
        ]
    else:
        direction = 1 if y2 >= y1 else -1
        points = [
            (x2, y2),
            (x2 - head // 2, y2 - direction * head),
            (x2 + head // 2, y2 - direction * head),
        ]
    draw.polygon(points, fill=fill)


def place_edge_label(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    text: str,
    *,
    font: ImageFont.ImageFont = FONT_SMALL,
    fill: str = "#0f172a",
    background: str = "#ffffff",
) -> None:
    width, height = text_size(draw, text, font)
    x, y = center
    pad = 10
    box = (x - width // 2 - pad, y - height // 2 - pad, x + width // 2 + pad, y + height // 2 + pad)
    draw.rounded_rectangle(box, radius=12, fill=background, outline="#cbd5e1", width=2)
    draw_centered_text(draw, box, text, font=font, fill=fill)


def generate_system_overview() -> None:
    image = Image.new("RGB", (2200, 1450), "#f8fafc")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Многосессионная архитектура backend", font=FONT_TITLE, fill="#0f172a")

    rounded_box(draw, (120, 130, 430, 260), "Web UI /\nвкладки браузера", fill="#dbeafe")
    rounded_box(draw, (640, 130, 1080, 260), "RuntimeSessionManager /\nсерверный контур", fill="#c7d2fe")

    draw.rounded_rectangle((120, 450, 560, 1030), radius=26, fill="#fee2e2", outline="#7f1d1d", width=5)
    draw.text((195, 490), "Сессии real-robot", font=FONT_BOLD, fill="#7f1d1d")
    rounded_box(draw, (185, 600, 500, 735), "SessionKey:\n(clientId, real-robot)", fill="#fff1f2", outline="#7f1d1d")
    rounded_box(draw, (185, 800, 500, 935), "TCP `5051` + камера +\nbridge сенсоров", fill="#fff1f2", outline="#7f1d1d")

    draw.rounded_rectangle((740, 450, 1180, 1030), radius=26, fill="#dcfce7", outline="#166534", width=5)
    draw.text((800, 490), "Сцены Unity", font=FONT_BOLD, fill="#14532d")
    rounded_box(draw, (805, 590, 1115, 725), "Unity world:\n(host, port)", fill="#f0fdf4", outline="#166534")
    rounded_box(draw, (805, 790, 1115, 935), "clientId -> world binding\nSignalR groups", fill="#f0fdf4", outline="#166534")

    rounded_box(draw, (1400, 240, 2000, 420), "Реальный стенд KS0223", fill="#fecaca")
    draw.rounded_rectangle((1400, 630, 2000, 980), radius=26, fill="#dcfce7", outline="#166534", width=5)
    draw.text((1535, 660), "Общая сцена Unity", font=FONT_BOLD, fill="#14532d")
    rounded_box(draw, (1510, 740, 1890, 860), "track / agents /\ncamera mode / reset", fill="#ecfccb", outline="#166534")

    arrow(draw, (430, 195), (640, 195))
    arrow(draw, (860, 260), (860, 450))
    arrow(draw, (820, 260), (360, 450))
    arrow(draw, (560, 720), (1400, 330))
    arrow(draw, (1180, 720), (1400, 720))

    place_edge_label(draw, (530, 155), "clientId + runtimeMode")
    place_edge_label(draw, (1110, 675), "unity-sim")
    place_edge_label(draw, (980, 330), "маршрутизация сессий")
    place_edge_label(draw, (980, 555), "привязка к миру")
    place_edge_label(draw, (980, 805), "общая сцена")
    place_edge_label(draw, (970, 260), "состояние / камера / события")

    rounded_box(
        draw,
        (1260, 1060, 2040, 1280),
        "Одна вкладка адресуется через clientId и runtimeMode.\nДля Unity несколько вкладок могут быть связаны\nс одним endpoint и работать в общей сцене.",
        fill="#ffffff",
        outline="#94a3b8",
        font=FONT_SMALL,
    )
    image.save(SYSTEM_OVERVIEW_PNG)


def generate_system_interaction() -> None:
    image = Image.new("RGB", (2200, 1520), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Разделение уровня мира и клиентской сессии", font=FONT_TITLE, fill="#0f172a")

    draw.rounded_rectangle((120, 200, 1020, 1320), radius=26, fill="#dcfce7", outline="#166534", width=5)
    draw.text((180, 240), "Уровень мира симуляции", font=FONT_BOLD, fill="#14532d")
    rounded_box(draw, (220, 360, 920, 500), "Unity world: host:port", fill="#f0fdf4", outline="#166534")
    rounded_box(draw, (220, 590, 920, 740), "track / agents / camera mode", fill="#ecfccb", outline="#166534")
    rounded_box(draw, (220, 840, 920, 990), "reset / runtime-selection", fill="#ecfccb", outline="#166534")
    rounded_box(draw, (220, 1080, 920, 1225), "общая сцена и видимость агентов", fill="#f0fdf4", outline="#166534")

    draw.rounded_rectangle((1180, 200, 2060, 1320), radius=26, fill="#dbeafe", outline="#1d4ed8", width=5)
    draw.text((1240, 240), "Уровень клиентской сессии", font=FONT_BOLD, fill="#1d4ed8")
    rounded_box(draw, (1280, 360, 1960, 500), "Client session:\nclientId + runtimeMode", fill="#eff6ff", outline="#1d4ed8")
    rounded_box(draw, (1280, 590, 1960, 740), "control agent / camera agent", fill="#dbeafe", outline="#1d4ed8")
    rounded_box(draw, (1280, 840, 1960, 990), "camera / status / sensors / command", fill="#dbeafe", outline="#1d4ed8")
    rounded_box(draw, (1280, 1080, 1960, 1225), "client-selection / BindClient", fill="#eff6ff", outline="#1d4ed8")

    arrow(draw, (920, 430), (1280, 430))
    arrow(draw, (920, 660), (1280, 660))
    arrow(draw, (920, 915), (1280, 915))
    place_edge_label(draw, (1100, 390), "общая конфигурация")
    place_edge_label(draw, (1100, 620), "адресная привязка")
    place_edge_label(draw, (1100, 875), "данные по сессии")

    draw.rounded_rectangle((360, 1360, 1820, 1460), radius=20, fill="#ffffff", outline="#94a3b8", width=4)
    draw.text(
        (420, 1390),
        "Уровень мира изменяет общую сцену. Уровень клиентской сессии определяет, как конкретная вкладка работает внутри этой сцены.",
        font=FONT_SMALL,
        fill="#111827",
    )
    image.save(SYSTEM_INTERACTION_PNG)


def generate_api_contract_diagram() -> None:
    image = Image.new("RGB", (2200, 1450), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Схема адресации runtime API", font=FONT_TITLE, fill="#0f172a")
    rounded_box(draw, (120, 170, 2080, 300), "Операции адресуются через clientId и runtimeMode", fill="#dbeafe")

    columns = [
        (150, 420, 640, 1160, "#fee2e2", "#7f1d1d", "Подключение и состояние", ["POST /api/connection/connect", "POST /api/connection/disconnect", "GET /api/status", "GET /api/health"]),
        (720, 420, 1210, 1160, "#dcfce7", "#166534", "Управление и данные", ["POST /api/command", "GET /api/camera/*", "GET /api/sensors/*", "POST /api/sensors/*"]),
        (1290, 420, 1780, 1160, "#ede9fe", "#6d28d9", "Уровень мира", ["GET /api/unity/runtime-catalog", "POST /api/unity/runtime-selection", "track / agents / reset"]),
        (1860, 420, 2080, 1160, "#eff6ff", "#1d4ed8", "Уровень сессии", ["POST /api/unity/client-selection", "BindClient(clientId)", "control agent", "camera agent"]),
    ]
    for x1, y1, x2, y2, fill, outline, title, lines in columns:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline=outline, width=4)
        draw.text((x1 + 25, y1 + 20), title, font=FONT_BOLD, fill=outline)
        yy = y1 + 120
        for line in lines:
            draw.text((x1 + 25, yy), f"• {line}", font=FONT_SMALL, fill="#111827")
            yy += 115
    image.save(API_CONTRACT_PNG)


def generate_connection_ui_diagram() -> None:
    image = Image.new("RGB", (2200, 1450), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Схема экрана подключения", font=FONT_TITLE, fill="#0f172a")

    draw.rounded_rectangle((180, 170, 2020, 1280), radius=30, fill="#101a24", outline="#334155", width=5)
    draw.text((260, 230), "Подключение", font=FONT_BOLD, fill="#f8fafc")
    rounded_box(draw, (280, 340, 840, 470), "Configured endpoint\n127.0.0.1:8000", fill="#dbeafe")
    rounded_box(draw, (1040, 340, 1600, 470), "Connected endpoint\n127.0.0.1:8000", fill="#dcfce7")
    rounded_box(draw, (280, 590, 760, 710), "Runtime mode\nunity-sim", fill="#eff6ff")
    rounded_box(draw, (840, 590, 1220, 710), "Host\n127.0.0.1", fill="#eff6ff")
    rounded_box(draw, (1300, 590, 1600, 710), "Port\n8000", fill="#eff6ff")
    rounded_box(draw, (280, 860, 700, 980), "Connect / Disconnect", fill="#ecfccb")
    rounded_box(draw, (780, 860, 1220, 980), "Status chips /\nlatency / errors", fill="#ede9fe")
    rounded_box(draw, (1300, 860, 1800, 980), "Unity scene selector", fill="#dbeafe")
    image.save(CONNECTION_UI_PNG)


def generate_unity_popup_diagram() -> None:
    image = Image.new("RGB", (2200, 1450), "#f8fafc")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Схема настроек Unity", font=FONT_TITLE, fill="#0f172a")

    draw.rounded_rectangle((420, 160, 1780, 1260), radius=28, fill="#ffffff", outline="#334155", width=5)
    draw.text((520, 220), "Настройки Unity симулятора", font=FONT_BOLD, fill="#111827")
    rounded_box(draw, (540, 350, 1060, 470), "Трек\ntrack.roadsystem_realistic", fill="#dcfce7")
    rounded_box(draw, (1140, 350, 1660, 470), "Camera mode\nspectator", fill="#dcfce7")
    rounded_box(draw, (540, 560, 1660, 760), "Список машинок на трассе\n• ego / vehicle.arcade.blue.v1\n• npc-red / vehicle.arcade.red.v1", fill="#dbeafe")
    rounded_box(draw, (540, 840, 1060, 960), "Control agent\nego", fill="#eff6ff")
    rounded_box(draw, (1140, 840, 1660, 960), "Camera agent\nnpc-red", fill="#eff6ff")
    rounded_box(draw, (540, 1060, 1660, 1160), "Применение конфигурации мира\nчерез runtime-selection", fill="#ede9fe")
    image.save(UNITY_POPUP_PNG)


def generate_shared_world_tabs_diagram() -> None:
    image = Image.new("RGB", (2200, 1450), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Работа двух вкладок в одной сцене Unity", font=FONT_TITLE, fill="#0f172a")

    rounded_box(draw, (150, 240, 520, 390), "Вкладка A\nclientId = tab-a\ncontrol = ego\ncamera = ego", fill="#dbeafe")
    rounded_box(draw, (150, 760, 520, 910), "Вкладка B\nclientId = tab-b\ncontrol = npc-red\ncamera = npc-red", fill="#dbeafe")

    draw.rounded_rectangle((760, 160, 2050, 1180), radius=30, fill="#dcfce7", outline="#166534", width=5)
    draw.text((840, 210), "Unity world: 127.0.0.1:8000", font=FONT_BOLD, fill="#14532d")
    rounded_box(draw, (960, 390, 1390, 540), "Agent ego", fill="#ecfccb", outline="#166534")
    rounded_box(draw, (1510, 700, 1940, 850), "Agent npc-red", fill="#ecfccb", outline="#166534")
    rounded_box(draw, (980, 900, 1910, 1070), "Общие track / physics /\nreset / camera mode", fill="#f0fdf4", outline="#166534")

    arrow(draw, (520, 315), (960, 465))
    arrow(draw, (520, 835), (1510, 775))
    place_edge_label(draw, (730, 360), "команды и камера")
    place_edge_label(draw, (980, 800), "команды и камера")
    image.save(SHARED_WORLD_TABS_PNG)


def generate_parallel_runtime_diagram() -> None:
    image = Image.new("RGB", (2200, 1450), "#f8fafc")
    draw = ImageDraw.Draw(image)
    draw.text((70, 50), "Параллельная работа с несколькими средами исполнения", font=FONT_TITLE, fill="#0f172a")

    rounded_box(draw, (100, 250, 520, 420), "Вкладка A\nunity local\n127.0.0.1:8000", fill="#dbeafe")
    rounded_box(draw, (100, 610, 520, 780), "Вкладка B\nunity remote\n192.168.1.50:8000", fill="#dbeafe")
    rounded_box(draw, (100, 970, 520, 1140), "Вкладка C\nreal-robot\n192.168.1.121:5051", fill="#dbeafe")
    rounded_box(draw, (770, 560, 1250, 760), "RuntimeSessionManager", fill="#c7d2fe")
    rounded_box(draw, (1510, 200, 2010, 420), "Unity world A\nlocal scene", fill="#dcfce7")
    rounded_box(draw, (1510, 560, 2010, 780), "Unity world B\nremote scene", fill="#dcfce7")
    rounded_box(draw, (1510, 920, 2010, 1140), "Реальный KS0223", fill="#fecaca")

    arrow(draw, (520, 335), (770, 620))
    arrow(draw, (520, 695), (770, 660))
    arrow(draw, (520, 1055), (770, 700))
    arrow(draw, (1250, 620), (1510, 310))
    arrow(draw, (1250, 660), (1510, 670))
    arrow(draw, (1250, 700), (1510, 1030))
    place_edge_label(draw, (645, 500), "маршрутизация сессий")
    place_edge_label(draw, (1380, 620), "выбор среды по параметрам")
    image.save(PARALLEL_RUNTIME_PNG)


def generate_diagrams() -> None:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    generate_system_overview()
    generate_system_interaction()
    generate_api_contract_diagram()
    generate_connection_ui_diagram()
    generate_unity_popup_diagram()
    generate_shared_world_tabs_diagram()
    generate_parallel_runtime_diagram()


def apply_main_heading_style(paragraph) -> None:
    paragraph.style = "Heading 1"
    paragraph.alignment = None
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.paragraph_format.line_spacing = None


def apply_structural_heading_style(paragraph) -> None:
    paragraph.style = "Heading 3"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.paragraph_format.line_spacing = None


def apply_normal_style(paragraph) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.paragraph_format.line_spacing = None


def apply_list_style(paragraph) -> None:
    paragraph.style = "List Paragraph"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.space_before = None
    paragraph.paragraph_format.space_after = None
    paragraph.paragraph_format.line_spacing = None


def apply_figure_caption_style(paragraph) -> None:
    paragraph.style = "Рисунок"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)


def apply_listing_title_style(paragraph) -> None:
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def apply_code_block_style(paragraph) -> None:
    paragraph.style = "Листинг"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def append_paragraph(doc: Document, text: str, *, style: str = "Normal"):
    paragraph = doc.add_paragraph()
    if text:
        paragraph.add_run(text)
    if style == "Heading 1":
        apply_main_heading_style(paragraph)
    elif style == "Heading 3":
        apply_structural_heading_style(paragraph)
    elif style == "List Paragraph":
        apply_list_style(paragraph)
    elif style == "Рисунок":
        apply_figure_caption_style(paragraph)
    elif style == "Листинг":
        apply_code_block_style(paragraph)
    else:
        apply_normal_style(paragraph)
    return paragraph


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)


def append_figure(doc: Document, image_path: Path, caption: str, *, width_cm: float = 16.0) -> None:
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Pt(0)
    image_paragraph.paragraph_format.left_indent = Pt(0)
    run = image_paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))

    caption_paragraph = doc.add_paragraph(caption)
    apply_figure_caption_style(caption_paragraph)


def style_table(table) -> None:
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.style = "Normal"
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.left_indent = Pt(0)
                paragraph.paragraph_format.right_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)


def append_table(doc: Document, caption: str, rows: list[list[str]]) -> None:
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.style = "Normal"
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_paragraph.paragraph_format.first_line_indent = Pt(0)
    caption_paragraph.paragraph_format.left_indent = Pt(0)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            table.cell(r_idx, c_idx).text = value
    style_table(table)


def append_listing(doc: Document, number: int, title: str, code: str) -> None:
    title_paragraph = doc.add_paragraph(f"Листинг {number} - {title}")
    apply_listing_title_style(title_paragraph)
    code_paragraph = doc.add_paragraph(code)
    apply_code_block_style(code_paragraph)


def clear_body_after_title_page(doc: Document) -> None:
    intro_paragraph = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "ВВЕДЕНИЕ")
    body = doc._body._element
    children = list(body)
    cutoff = children.index(intro_paragraph._element)
    end = len(children)
    if children and children[-1].tag.endswith("}sectPr"):
        end -= 1
    for child in children[cutoff:end]:
        body.remove(child)


def restore_title_page_format(doc: Document) -> None:
    title_indexes = [0, 1, 3, 4, 5, 6, 18, 19, 20, 21, 22, 23, 32]
    for idx in title_indexes:
        paragraph = doc.paragraphs[idx]
        paragraph.style = "Normal"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0


def build_report(doc: Document) -> None:
    append_paragraph(doc, "ВВЕДЕНИЕ", style="Heading 3")
    for paragraph in INTRO_TEXT:
        append_paragraph(doc, paragraph)

    append_paragraph(doc, "ЦЕЛЬ, ЗАДАЧИ И ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ ПРАКТИКИ", style="Heading 1")
    for paragraph in GOAL_TEXT:
        append_paragraph(doc, paragraph)
    append_table(doc, "Таблица 1 - Соответствие задач практики и полученных результатов", TABLE_1_ROWS)

    append_paragraph(doc, "КРАТКАЯ ХАРАКТЕРИСТИКА ОБЪЕКТА ПРАКТИКИ И ИСХОДНЫХ УСЛОВИЙ", style="Heading 1")
    for index, paragraph in enumerate(OBJECT_TEXT):
        append_paragraph(doc, paragraph)
        if index == 0:
            append_figure(doc, ROBOT_PNG, "Рисунок 1 - Роботизированная платформа Keyestudio KS0223", width_cm=11.5)

    append_paragraph(doc, "ЭТАПЫ ВЫПОЛНЕНИЯ ПРАКТИКИ И ЛИЧНЫЙ ВКЛАД СТУДЕНТА", style="Heading 1")
    for paragraph in STAGES_INTRO:
        append_paragraph(doc, paragraph)
    for number, item in enumerate(STAGE_ITEMS, start=1):
        append_paragraph(doc, f"{number}. {item}", style="List Paragraph")
    for paragraph in STAGES_OUTRO:
        append_paragraph(doc, paragraph)

    append_paragraph(doc, "ИСПОЛЬЗОВАННЫЕ ТЕХНОЛОГИИ И ИНСТРУМЕНТЫ", style="Heading 1")
    for paragraph in TECH_TEXT:
        append_paragraph(doc, paragraph)

    append_paragraph(doc, "ТЕХНИЧЕСКАЯ РЕАЛИЗАЦИЯ И ПОЛУЧЕННЫЕ РЕЗУЛЬТАТЫ", style="Heading 1")
    append_paragraph(doc, TECH_IMPL_TEXT[0])
    append_figure(doc, ARCHITECTURE_PNG, "Рисунок 2 - Схема продуктовой архитектуры единого WebUI-драйвера для KS0223")
    append_figure(doc, SEQUENCE_PNG, "Рисунок 3 - Диаграмма последовательности подключения оператора и отправки команд")
    for paragraph in TECH_IMPL_TEXT[1:4]:
        append_paragraph(doc, paragraph)
    append_paragraph(doc, TECH_IMPL_TEXT[4])
    append_figure(doc, CONTROL_UI_PNG, "Рисунок 4 - Основной экран WebUI-драйвера в режиме управления платформой KS0223")
    append_figure(doc, LOGS_UI_PNG, "Рисунок 5 - Экран журналирования, диагностики и управления логами")
    append_figure(doc, SHARED_WORLD_TABS_PNG, "Рисунок 6 - Схема работы двух клиентских сессий в общей сцене Unity")
    append_paragraph(doc, TECH_IMPL_TEXT[5])

    append_paragraph(doc, "ПРОВЕРКА РАБОТОСПОСОБНОСТИ И АНАЛИЗ РЕЗУЛЬТАТОВ", style="Heading 1")
    append_paragraph(doc, VERIFICATION_TEXT[0])
    append_paragraph(doc, VERIFICATION_TEXT[1])
    append_table(doc, "Таблица 2 - Проверка функциональных сценариев работы системы", TABLE_2_ROWS)
    append_paragraph(doc, VERIFICATION_TEXT[2])
    append_table(doc, "Таблица 3 - Результаты эксплуатационной проверки", TABLE_3_ROWS)
    append_paragraph(doc, VERIFICATION_TEXT[3])

    append_paragraph(doc, "ПРИОБРЕТЕННЫЕ НАВЫКИ И КОМПЕТЕНЦИИ", style="Heading 1")
    for paragraph in SKILLS_TEXT:
        append_paragraph(doc, paragraph)
    for number, item in enumerate(SKILL_ITEMS, start=1):
        append_paragraph(doc, f"{number}. {item}", style="List Paragraph")

    add_page_break(doc)
    append_paragraph(doc, "ЗАКЛЮЧЕНИЕ", style="Heading 3")
    for paragraph in CONCLUSION_TEXT:
        append_paragraph(doc, paragraph)

    add_page_break(doc)
    append_paragraph(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", style="Heading 3")
    for index, source in enumerate(SOURCES, start=1):
        paragraph = append_paragraph(doc, f"{index}. {source}")
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)

    add_page_break(doc)
    append_paragraph(doc, "ПРИЛОЖЕНИЕ А. ИСПОЛЬЗОВАННЫЕ ТЕХНОЛОГИИ И ИХ НАЗНАЧЕНИЕ", style="Heading 3")
    for paragraph in APPENDIX_A_TEXT:
        append_paragraph(doc, paragraph)

    add_page_break(doc)
    append_paragraph(doc, "ПРИЛОЖЕНИЕ Б. ДОПОЛНИТЕЛЬНЫЕ АРХИТЕКТУРНЫЕ СХЕМЫ", style="Heading 3")
    for paragraph in APPENDIX_B_TEXT:
        append_paragraph(doc, paragraph)
    append_figure(doc, SYSTEM_OVERVIEW_PNG, "Рисунок 7 - Многосессионная организация серверного контура и режимов работы")
    append_figure(doc, SYSTEM_INTERACTION_PNG, "Рисунок 8 - Разделение управления на уровне мира симуляции и клиентской сессии")
    append_figure(doc, CONNECTION_UI_PNG, "Рисунок 9 - Схема экрана подключения с параметрами Configured endpoint и Connected endpoint")
    append_figure(doc, UNITY_POPUP_PNG, "Рисунок 10 - Схема настройки сцены Unity и параметров клиентской сессии")
    append_figure(doc, API_CONTRACT_PNG, "Рисунок 11 - Схема адресации runtime API по параметрам clientId и runtimeMode")
    append_figure(doc, PARALLEL_RUNTIME_PNG, "Рисунок 12 - Параллельная работа с несколькими средами исполнения")

    add_page_break(doc)
    append_paragraph(doc, "ПРИЛОЖЕНИЕ В. ФРАГМЕНТЫ КОДА", style="Heading 3")
    append_paragraph(doc, "В приложении В приведены фрагменты кода, иллюстрирующие ключевые технические решения, реализованные в ходе практики.")
    append_listing(
        doc,
        1,
        "Хранение клиентских сессий и сцен Unity в RuntimeSessionManager",
        """private readonly ConcurrentDictionary<SessionKey, RealRuntimeSession> realSessions = new();
private readonly ConcurrentDictionary<UnityWorldKey, UnityWorldSession> unityWorlds = new();
private readonly ConcurrentDictionary<string, UnityClientBinding> unityClientBindings = new(StringComparer.Ordinal);

public static string GetClientGroup(string clientId) => $"client:{clientId}";""",
    )
    append_listing(
        doc,
        2,
        "Разделение настройки мира и клиентской сессии в backend API",
        """app.MapPost("/api/unity/runtime-selection", async (UnityRuntimeSelectionRequest request, RuntimeSessionManager runtimeSessionManager, CancellationToken cancellationToken) =>
{
    var catalog = await runtimeSessionManager.SetUnityRuntimeSelectionAsync(request, cancellationToken);
    return Results.Ok(catalog);
});

app.MapPost("/api/unity/client-selection", async (UnityClientSelectionRequest request, RuntimeSessionManager runtimeSessionManager, CancellationToken cancellationToken) =>
{
    var catalog = await runtimeSessionManager.SetUnityClientSelectionAsync(request, cancellationToken);
    return Results.Ok(catalog);
});""",
    )
    append_listing(
        doc,
        3,
        "Привязка SignalR-соединения к клиентской сессии",
        """public async Task BindClient(string clientId)
{
    var normalizedClientId = clientId.Trim();
    var group = RuntimeSessionManager.GetClientGroup(normalizedClientId);

    ConnectionBindings[Context.ConnectionId] = normalizedClientId;
    await Groups.AddToGroupAsync(Context.ConnectionId, group);
    await runtimeSessionManager.RegisterBoundClientAsync(normalizedClientId, Context.ConnectionId);
}""",
    )
    append_listing(
        doc,
        4,
        "Формирование конфигурации общей сцены Unity",
        """var payload = new
{
    selectedTrackId,
    selectedVehicleId,
    vehicleParams = new object[] { new { key = "camera.mode", value = cameraMode } },
    flags = new object[] { new { key = "agents.allow_empty", value = agentsSnapshot.Count == 0 ? "true" : "false" } },
    agents = agentsSnapshot.Select((agent, index) => new
    {
        agentId = string.IsNullOrWhiteSpace(agent.AgentId) ? $"agent-{index + 1}" : agent.AgentId,
        vehicleId = agent.VehicleId,
        isPrimary = agent.IsPrimary || index == 0,
    }).ToArray(),
};""",
    )

    add_page_break(doc)
    append_paragraph(doc, "ПРИЛОЖЕНИЕ Г. ДОПОЛНИТЕЛЬНЫЕ МАТЕРИАЛЫ И НАПРАВЛЕНИЯ РАЗВИТИЯ", style="Heading 3")
    for paragraph in APPENDIX_D_TEXT:
        append_paragraph(doc, paragraph)
    append_figure(doc, KIT_LIST_PNG, "Рисунок 13 - Пример состава набора KS0223 по документации Keyestudio", width_cm=13.5)


def main() -> None:
    generate_diagrams()
    doc = Document(REPORT_PATH)
    clear_body_after_title_page(doc)
    build_report(doc)
    restore_title_page_format(doc)
    doc.save(REPORT_PATH)


if __name__ == "__main__":
    main()
