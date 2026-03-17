#!/usr/bin/env python3
"""Build a standalone literature review DOCX for the practice topic."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[3]
OUTPUT_DOCX = ROOT / "docs/report/practice/literature_review.docx"


TITLE = "Обзор литературы по поставленной задаче"

INTRO_PARAGRAPHS = [
    (
        "Разработка программных контуров для робототехнических систем все чаще "
        "опирается на совместное использование физического стенда и симулятора. "
        "Такой подход позволяет безопасно отрабатывать алгоритмы управления, "
        "проверять операторские сценарии и переносить результаты из виртуальной "
        "среды на реальное устройство [1], [2]."
    ),
    (
        "Для рассматриваемой задачи это особенно важно, поскольку требуется "
        "создать единую программную среду управления и мониторинга для двух "
        "режимов работы: физического робототехнического стенда Keyestudio KS0223 "
        "и Unity-симулятора. Литература показывает, что наибольшую практическую "
        "ценность имеют решения, где объединены не только управление, но и "
        "видеоданные, сенсорная информация и средства диагностики [3], [5]."
    ),
    (
        "Следовательно, обзор литературы должен охватывать не только исследования "
        "по sim-to-real переносу, но и работы по цифровым двойникам, средствам "
        "моделирования в Unity, а также подходам к web-управлению роботами. "
        "Именно сочетание этих направлений образует теоретическую основу проекта "
        "единой web-среды управления и наблюдения для KS0223 и Unity [4], [6], [8]."
    ),
]

SECTION_2_PARAGRAPHS = [
    (
        "В обзорах по sim-to-real переносу подчеркивается, что обучение и "
        "отладка в симуляции позволяют снизить стоимость экспериментов, сократить "
        "риск повреждения оборудования и быстрее проверять гипотезы по "
        "управлению роботами [1]. Однако такой подход сталкивается с так "
        "называемым разрывом между моделью и реальностью, когда поведение "
        "алгоритма в виртуальной среде не полностью совпадает с поведением на "
        "физическом стенде."
    ),
    (
        "Одним из наиболее известных путей уменьшения этого разрыва стала "
        "рандомизация параметров среды, описанная в работе J. Tobin и соавторов [2]. "
        "Смысл подхода состоит в том, чтобы во время симуляции варьировать "
        "визуальные и физические характеристики сцены, делая алгоритм более "
        "устойчивым к изменениям при переходе на реальное устройство. Для "
        "прикладной задачи управления KS0223 этот вывод важен тем, что сам "
        "симулятор должен рассматриваться не как изолированная демонстрация, а "
        "как среда, близкая по интерфейсам наблюдения и управления к реальному "
        "стенду."
    ),
    (
        "Литература по цифровым двойникам расширяет этот взгляд. В работе "
        "J. A. Douthwaite и соавторов цифровой двойник рассматривается как "
        "синхронизированная программная модель робототехнической системы, "
        "пригодная для анализа безопасности, наблюдения и тестирования [3]. "
        "Похожая идея прослеживается в работе TELESIM, где цифровой двойник "
        "используется как промежуточный слой для телеуправления роботизированной "
        "системой и наблюдения за ее состоянием [4]. Для рассматриваемого проекта "
        "это означает, что Unity-симулятор должен выполнять роль не отдельной "
        "учебной сцены, а цифровой модели физической платформы, связанной с "
        "общим пользовательским сценарием."
    ),
]

SECTION_3_PARAGRAPHS = [
    (
        "Практическая реализация цифрового двойника требует среды, способной "
        "поддерживать физическое моделирование, работу камер, агентов и "
        "внешних программных интерфейсов. Unity и экосистема ML-Agents "
        "широко применяются именно в таких задачах, поскольку позволяют "
        "создавать управляемые сцены, связывать логику агентов с сенсорами и "
        "подключать внешние алгоритмы обучения и управления [6], [7]."
    ),
    (
        "Для целей рассматриваемой работы Unity важен не только как движок "
        "трехмерной визуализации. Важнее то, что он поддерживает сценарное "
        "моделирование, конфигурацию камеры, смену активных объектов и "
        "наблюдение за состоянием сцены. Это делает его подходящей основой "
        "для симулятора, который должен быть встроен в общий контур управления "
        "и использоваться по тем же прикладным принципам, что и физический "
        "стенд."
    ),
    (
        "Таким образом, публикации и официальная документация по Unity "
        "подтверждают, что данная среда подходит для построения цифрового "
        "двойника робототехнической платформы, в котором можно объединить "
        "визуальную модель, каналы наблюдения и интерфейсы управления. Для "
        "поставленной задачи это особенно существенно, поскольку симулятор "
        "должен быть не самостоятельным продуктом, а частью общей среды "
        "для sim-to-real исследований."
    ),
]

SECTION_4_PARAGRAPHS = [
    (
        "Отдельное направление литературы посвящено телеуправлению и "
        "операторским интерфейсам роботов. В обзоре M. D. Moniruzzaman и "
        "соавторов подчеркивается, что качественная система телеуправления "
        "должна обеспечивать оператору не только передачу команд, но и "
        "достаточный объем обратной связи: видео, телеметрию, диагностические "
        "данные и информацию о состоянии канала связи [5]. Это напрямую "
        "соотносится с рассматриваемой задачей создания единого WebUI для "
        "стенда и симулятора."
    ),
    (
        "Для реализации подобных интерфейсов важную роль играют web-технологии. "
        "Проект Robot Web Tools показывает, что браузер может выступать "
        "полноценной средой наблюдения и управления роботами, если между "
        "робототехническим контуром и пользовательским интерфейсом существует "
        "понятный программный слой обмена данными [8]. Документация ROS 2, "
        "в свою очередь, демонстрирует важность модульной архитектуры и "
        "разделения транспорта данных, прикладной логики и пользовательских "
        "средств мониторинга [9]."
    ),
    (
        "Для проекта с KS0223 этот вывод означает, что web-интерфейс не "
        "может ограничиваться одной страницей с кнопками движения. Он должен "
        "объединять подключение, передачу команд, просмотр видеоданных, "
        "получение сенсорной информации и диагностику состояния системы. "
        "Аппаратная документация KS0223 подтверждает наличие нескольких "
        "разнородных подсистем, включая камеру, ультразвуковой датчик, "
        "сервопривод и линейные сенсоры [10]. Поэтому задача объединения этих "
        "каналов в одной программной среде является содержательно обоснованной, "
        "а не искусственно сформулированной."
    ),
]

CONCLUSION_PARAGRAPHS = [
    (
        "Проведенный обзор показывает, что для поставленной задачи "
        "существует устойчивая теоретическая и прикладная база. Исследования "
        "по sim-to-real переносу подтверждают важность тесной связи между "
        "симуляцией и реальным стендом [1], [2], а работы по цифровым "
        "двойникам показывают ценность синхронизированной виртуальной модели "
        "для наблюдения, проверки и безопасной отладки робототехнических "
        "систем [3], [4]."
    ),
    (
        "Официальная документация Unity, ROS 2, Robot Web Tools и KS0223 "
        "подтверждает практическую реализуемость такого подхода [6]–[10]. "
        "Следовательно, разработка единой программной среды управления, "
        "наблюдения и диагностики для KS0223 и Unity-симулятора является "
        "актуальной и обоснованной задачей. Литература поддерживает именно "
        "тот вектор, который используется в проекте: физический стенд и "
        "симулятор должны работать в рамках одного пользовательского сценария, "
        "что создает основу для дальнейших sim-to-real исследований."
    ),
]

TABLE_ROWS = [
    ["Группа источников", "Что рассматривается", "Значение для поставленной задачи"],
    [
        "Sim-to-real перенос [1], [2]",
        "Перенос алгоритмов и сценариев из симуляции на физический робот, проблема reality gap",
        "Обосновывает необходимость единого контура управления и сопоставимых интерфейсов для стенда и симулятора",
    ],
    [
        "Цифровые двойники [3], [4]",
        "Синхронизированная виртуальная модель, наблюдение, безопасность, teleoperation",
        "Подтверждает роль Unity-симулятора как цифрового двойника, а не отдельной демонстрационной сцены",
    ],
    [
        "Средства моделирования Unity [6], [7]",
        "Сцены, агенты, камеры, взаимодействие с внешними алгоритмами и API",
        "Показывает пригодность Unity для моделирования операторских и исследовательских сценариев",
    ],
    [
        "Web и robotics middleware [5], [8], [9]",
        "Телеуправление, browser-интерфейсы, модульный обмен данными",
        "Поддерживает идею единого web-интерфейса для команд, видео, сенсоров и диагностики",
    ],
    [
        "Аппаратная документация KS0223 [10]",
        "Состав платформы, сенсоры, камера и исполнительные узлы",
        "Связывает литературный обзор с конкретной физической платформой проекта",
    ],
]

SOURCES = [
    "1. Zhu W., Guo X., Owaki D., Kutsuzawa K., Hayashibe M. A Survey of Sim-to-Real Transfer Techniques Applied to Reinforcement Learning for Bioinspired Robots // IEEE Transactions on Neural Networks and Learning Systems. 2023. Vol. 34, No. 7. P. 3444-3459. DOI: 10.1109/TNNLS.2021.3112718.",
    "2. Tobin J., Fong R., Ray A., Schneider J., Zaremba W., Abbeel P. Domain Randomization for Transferring Deep Neural Networks from Simulation to the Real World // arXiv preprint arXiv:1703.06907. 2017. URL: https://arxiv.org/abs/1703.06907 (дата обращения: 17.03.2026).",
    "3. Douthwaite J. A., Lesage B., Gleirscher M., Calinescu R., Aitken J. M., Alexander R. A Modular Digital Twinning Framework for Safety Assurance of Collaborative Robotics // Frontiers in Robotics and AI. 2021. Vol. 8. DOI: 10.3389/frobt.2021.758099.",
    "4. Audonnet F. P., Grizou J., Hamilton A., Aragon-Camarasa G. TELESIM: A Modular and Plug-and-Play Framework for Robotic Arm Teleoperation using a Digital Twin // arXiv preprint arXiv:2309.10579. 2023. URL: https://arxiv.org/abs/2309.10579 (дата обращения: 17.03.2026).",
    "5. Moniruzzaman M. D., Rassau A., Chai D., Islam S. M. S. Teleoperation methods and enhancement techniques for mobile robots: A comprehensive survey // Robotics and Autonomous Systems. 2022. Vol. 150. Art. 103973. DOI: 10.1016/j.robot.2021.103973.",
    "6. Unity Technologies. Unity 6.1 User Manual [Электронный ресурс] // Unity Documentation : [сайт]. URL: https://docs.unity3d.com/6000.1/Documentation/Manual/ (дата обращения: 17.03.2026).",
    "7. Unity Technologies. ML Agents [Электронный ресурс] // Unity Documentation : [сайт]. URL: https://docs.unity3d.com/6000.1/Documentation/Manual/com.unity.ml-agents.html (дата обращения: 17.03.2026).",
    "8. Robot Web Tools [Электронный ресурс] // Robot Web Tools : [сайт]. URL: https://robotwebtools.github.io/ (дата обращения: 17.03.2026).",
    "9. Open Robotics. ROS 2 Documentation: Humble [Электронный ресурс] // ROS 2 Documentation : [сайт]. URL: https://docs.ros.org/en/humble/index.html (дата обращения: 17.03.2026).",
    "10. Keyestudio. Raspberry Pi Smart Car Documentation [Электронный ресурс] // Keyestudio Docs : [сайт]. URL: https://docs.keyestudio.com/projects/KS0223/en/latest/ (дата обращения: 17.03.2026).",
]


def set_page_layout(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)


def configure_normal_style(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(14)


def set_paragraph_format(paragraph, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.first_line_indent = Cm(1.25) if align == WD_ALIGN_PARAGRAPH.JUSTIFY else Cm(0)
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(14)
        run.bold = bold


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    set_paragraph_format(paragraph)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(6)
    fmt.first_line_indent = Cm(0)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)
    run.bold = True


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(12)
    fmt.first_line_indent = Cm(0)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(16)
    run.bold = True


def add_table(document: Document, rows: list[list[str]]) -> None:
    caption = document.add_paragraph("Таблица 1 - Основные группы источников и их значение для задачи")
    set_paragraph_format(caption, align=WD_ALIGN_PARAGRAPH.LEFT)
    caption.paragraph_format.first_line_indent = Cm(0)

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(4.3), Cm(5.7), Cm(7.0)]

    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
                fmt = paragraph.paragraph_format
                fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                fmt.first_line_indent = Cm(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(12)
                    if row_idx == 0:
                        run.bold = True

    document.add_paragraph()


def add_sources(document: Document, sources: list[str]) -> None:
    add_heading(document, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for source in sources:
        paragraph = document.add_paragraph(source)
        set_paragraph_format(paragraph)
        paragraph.paragraph_format.first_line_indent = Cm(0)


def enable_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)


def build_document() -> Document:
    document = Document()
    set_page_layout(document)
    configure_normal_style(document)
    enable_update_fields_on_open(document)

    add_title(document, TITLE)

    add_heading(document, "1 Введение и постановка задачи")
    for paragraph in INTRO_PARAGRAPHS:
        add_body_paragraph(document, paragraph)

    add_table(document, TABLE_ROWS)

    add_heading(document, "2 Sim-to-real подход и цифровые двойники в робототехнике")
    for paragraph in SECTION_2_PARAGRAPHS:
        add_body_paragraph(document, paragraph)

    add_heading(document, "3 Unity как среда моделирования робототехнических систем")
    for paragraph in SECTION_3_PARAGRAPHS:
        add_body_paragraph(document, paragraph)

    add_heading(document, "4 Web-интерфейсы и программные контуры управления роботами")
    for paragraph in SECTION_4_PARAGRAPHS:
        add_body_paragraph(document, paragraph)

    add_heading(document, "5 Вывод")
    for paragraph in CONCLUSION_PARAGRAPHS:
        add_body_paragraph(document, paragraph)

    add_sources(document, SOURCES)
    return document


def main() -> None:
    document = build_document()
    document.save(OUTPUT_DOCX)
    print(f"Saved {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
